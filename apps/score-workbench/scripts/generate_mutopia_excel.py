import os
import re
import sys
import time
from fractions import Fraction
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding='utf-8')

MUTOPIA_DIR = r'K:\MusicSearch\data\classical-corpora\mutopia-project'
FTP_DIR = os.path.join(MUTOPIA_DIR, 'ftp')
OUTPUT_XLSX = r'K:\MusicSearch\data\Mutopia_Score_Catalog.xlsx'
OUTPUT_XLSX_ALT = os.path.join(MUTOPIA_DIR, 'Mutopia_Score_Catalog.xlsx')

def clean_markup(text):
    if not text:
        return ''
    # If text contains quoted strings inside, e.g. \markup { "Wolfgang Amadeus Mozart" }
    quoted = re.findall(r'["\']([^"\']{3,})["\']', text)
    if quoted and len(quoted[0]) > 3 and not quoted[0].startswith('\\'):
        text = ' '.join(quoted)
    text = re.sub(r'\\markup\b', '', text)
    text = re.sub(r'\\[a-zA-Z\-]+', '', text)
    text = re.sub(r'\b(center-align|right-align|left-align|column|line|italic|bold|small|large|huge|fontsize|font|raise|null|box|pad)\b', '', text, flags=re.IGNORECASE)
    text = re.sub(r'#[0-9\.\-]+', '', text)
    text = re.sub(r'[\"\'\{\}\\\(\)]', '', text)
    text = re.sub(r'^\s*-\s*', '', text)
    text = re.sub(r'\s+', ' ', text)
    return text.strip()

def get_field(content, field_name):
    # Match field = "..." or field = '...'
    m = re.search(r'^\s*' + field_name + r'\s*=\s*["\']([^"\']+)["\']', content, re.MULTILINE | re.IGNORECASE)
    if m:
        return m.group(1).strip()
    m2 = re.search(r'^\s*' + field_name + r'\s*=\s*([^\n\r%]+)', content, re.MULTILINE | re.IGNORECASE)
    if m2:
        val = m2.group(1).strip()
        return clean_markup(val)
    return ''

def calculate_measures(content):
    # Method 1: Bar number comments (% 120)
    bar_comments = re.findall(r'%\s*(\d+)\s*$', content, re.MULTILINE)
    if bar_comments:
        nums = [int(n) for n in bar_comments if n.isdigit() and 5 < int(n) < 5000]
        if nums:
            return max(nums), "주석 마디번호 (% N)"

    # Detect staves
    staves = (
        len(re.findall(r'\\context\s+Staff', content, re.IGNORECASE)) or
        len(re.findall(r'\\new\s+Staff', content, re.IGNORECASE)) or
        len(re.findall(r'\\relative', content, re.IGNORECASE)) or
        1
    )
    staves = max(1, min(12, staves))

    # Method 2: Bar check pipes (|)
    pipes = len(re.findall(r'\|', content))
    if pipes >= 8:
        est = round(pipes / staves)
        if est >= 4:
            return est, "바체크(|) 집계"

    # Method 3: Accumulate note durations
    time_m = re.search(r'\\time\s+(\d+)\/(\d+)', content)
    if time_m:
        bar_len = Fraction(int(time_m.group(1)), int(time_m.group(2)))
    else:
        bar_len = Fraction(4, 4)

    music_part = content
    header_idx = content.find(r'\score')
    if header_idx != -1:
        music_part = content[header_idx:]

    notes = re.findall(r'(?:[a-grsR]|>)(?:\'|,)*(\d+)(\.*)', music_part)
    if notes and len(notes) > 15:
        tot = Fraction(0, 1)
        for d, dots in notes:
            if not d.isdigit() or int(d) == 0:
                continue
            val = Fraction(1, int(d))
            dot_val = val
            for _ in range(len(dots)):
                dot_val /= 2
                val += dot_val
            tot += val
        meas = round(float(tot / bar_len) / staves)
        if meas >= 3:
            return meas, "음가 누적 환산"

    # Fallback: estimate from music line count
    music_lines = len(content.splitlines()) - 40
    if music_lines > 10:
        return max(4, round(music_lines * 1.5 / staves)), "코드 라인 기반 추정"

    return 0, "미확인"

def calculate_pages(content, measures, instrument):
    inst_lower = (instrument or '').lower()
    
    # 1. Explicit page-count
    pc_m = re.search(r'page-count\s*=\s*#?(\d+)', content, re.IGNORECASE)
    if pc_m:
        return int(pc_m.group(1)), "명시적 설정 (page-count)"

    # 2. \pageBreak
    page_breaks = len(re.findall(r'\\pageBreak', content))
    if page_breaks > 0:
        return page_breaks + 1, "페이지 브레이크 (\\pageBreak)"

    # 3. Density estimation
    if 'orchestra' in inst_lower or 'symphon' in inst_lower:
        rate = 6
    elif 'quartet' in inst_lower or 'quintet' in inst_lower or 'ensemble' in inst_lower:
        rate = 10
    elif 'choir' in inst_lower or 'chorus' in inst_lower:
        rate = 12
    elif 'piano' in inst_lower or 'harpsichord' in inst_lower or 'organ' in inst_lower or 'keyboard' in inst_lower:
        rate = 18
    elif 'guitar' in inst_lower or 'lute' in inst_lower:
        rate = 22
    elif 'violin' in inst_lower or 'cello' in inst_lower or 'flute' in inst_lower or 'voice' in inst_lower:
        rate = 24
    else:
        rate = 16

    est_pages = max(1, round(measures / rate)) if measures > 0 else 1
    return est_pages, "마디 밀도 추정"

def main():
    print(f"Scanning Mutopia files in {FTP_DIR}...")
    t0 = time.time()

    all_ly_files = []
    for root, dirs, files in os.walk(FTP_DIR):
        if '.git' in root:
            continue
        for f in files:
            if f.lower().endswith('.ly'):
                all_ly_files.append(os.path.join(root, f))

    all_ly_files.sort()
    print(f"Found {len(all_ly_files)} .ly files. Parsing metadata and metrics...")

    records = []
    composer_stats = {}

    for idx, fpath in enumerate(all_ly_files):
        rel_path = os.path.relpath(fpath, MUTOPIA_DIR).replace('\\', '/')
        file_size_kb = round(os.path.getsize(fpath) / 1024, 2)
        folder_path = os.path.dirname(fpath)
        
        # Calculate folder total size (including .ily files)
        folder_size_kb = 0
        ily_count = 0
        try:
            for item in os.listdir(folder_path):
                if item.lower().endswith('.ily'):
                    ily_count += 1
                item_full = os.path.join(folder_path, item)
                if os.path.isfile(item_full):
                    folder_size_kb += os.path.getsize(item_full)
            folder_size_kb = round(folder_size_kb / 1024, 2)
        except Exception:
            folder_size_kb = file_size_kb

        try:
            with open(fpath, 'r', encoding='utf-8', errors='ignore') as fp:
                content = fp.read()
        except Exception:
            content = ''

        # Header fields
        title = get_field(content, 'title') or get_field(content, 'mutopiatitle')
        subtitle = get_field(content, 'subtitle')
        piece = get_field(content, 'piece')
        composer = get_field(content, 'composer') or get_field(content, 'mutopiacomposer')
        composer_code = get_field(content, 'mutopiacomposer')
        opus = get_field(content, 'opus') or get_field(content, 'mutopiaopus')
        instrument = get_field(content, 'mutopiainstrument') or get_field(content, 'instrument')
        style = get_field(content, 'style')
        date = get_field(content, 'date')
        license_str = get_field(content, 'copyright') or get_field(content, 'license')
        maintainer = get_field(content, 'maintainer')
        footer = get_field(content, 'footer')
        ly_version = get_field(content, 'version')

        # Fallback for composer code & composer name from path: ftp/<ComposerCode>/...
        parts = rel_path.split('/')
        if len(parts) >= 2 and not composer_code:
            composer_code = parts[1]
        if not composer:
            composer = composer_code

        if not title:
            # Fallback to folder name or filename
            title = os.path.splitext(os.path.basename(fpath))[0]

        # Calculate measures & pages
        measures, m_method = calculate_measures(content)
        pages, p_method = calculate_pages(content, measures, instrument)

        rec = {
            'no': idx + 1,
            'footer': footer,
            'composer': clean_markup(composer),
            'composer_code': composer_code,
            'title': clean_markup(title),
            'piece_mvt': clean_markup(piece or subtitle),
            'opus': clean_markup(opus),
            'instrument': clean_markup(instrument),
            'style': clean_markup(style),
            'date': clean_markup(date),
            'measures': measures,
            'm_method': m_method,
            'pages': pages,
            'p_method': p_method,
            'size_kb': file_size_kb,
            'folder_size_kb': folder_size_kb,
            'ily_count': ily_count,
            'rel_path': rel_path,
            'filename': os.path.basename(fpath),
            'version': ly_version,
            'license': clean_markup(license_str),
            'maintainer': clean_markup(maintainer)
        }
        records.append(rec)

        # Aggregate composer stats by composer_code
        comp_key = composer_code or 'Unknown'
        if comp_key not in composer_stats:
            composer_stats[comp_key] = {
                'code': composer_code,
                'names': {},
                'count': 0,
                'total_measures': 0,
                'total_pages': 0,
                'total_size_kb': 0,
                'instruments': set(),
                'styles': set()
            }
        cs = composer_stats[comp_key]
        c_clean = clean_markup(composer)
        if c_clean and c_clean != composer_code:
            cs['names'][c_clean] = cs['names'].get(c_clean, 0) + 1
        cs['count'] += 1
        cs['total_measures'] += measures
        cs['total_pages'] += pages
        cs['total_size_kb'] += file_size_kb
        if clean_markup(instrument):
            cs['instruments'].add(clean_markup(instrument))
        if clean_markup(style):
            cs['styles'].add(clean_markup(style))

        if (idx + 1) % 500 == 0 or (idx + 1) == len(all_ly_files):
            print(f"Processed {idx + 1}/{len(all_ly_files)} files [{round(time.time() - t0, 1)}s]")

    print(f"Parsing finished in {round(time.time() - t0, 1)}s. Building Excel workbook with openpyxl...")

    # Build Excel Workbook
    wb = openpyxl.Workbook()

    # Sheet 1: Entire Scores
    ws1 = wb.active
    ws1.title = "Mutopia 악보 목록"
    ws1.views.sheetView[0].showGridLines = True

    headers = [
        "연번", "Mutopia 식별자 (Footer ID)", "작곡가", "작곡가 코드", "작품명 (Title)",
        "악장 / 소제목", "작품번호 (Opus/BWV/KV)", "편성 / 악기", "음악 양식", "작곡 연도",
        "마디 수 (Measures)", "마디 산출 기준", "쪽 수 (Pages)", "쪽수 산출 기준",
        "파일 용량 (KB)", "폴더 합계 용량 (KB)", "포함 성부보 (.ily) 수", "상대 경로",
        "파일명", "LilyPond 버전", "라이선스", "입력자 (Maintainer)"
    ]

    ws1.append(headers)

    # Styles
    font_header = Font(name="맑은 고딕", size=11, bold=True, color="FFFFFF")
    fill_header = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")
    border_thin = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    fill_zebra = PatternFill(start_color="F9FAFC", end_color="F9FAFC", fill_type="solid")
    font_data = Font(name="맑은 고딕", size=10)

    for col_num in range(1, len(headers) + 1):
        cell = ws1.cell(row=1, column=col_num)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
    ws1.row_dimensions[1].height = 28

    for r_idx, r in enumerate(records, start=2):
        row_values = [
            r['no'], r['footer'], r['composer'], r['composer_code'], r['title'],
            r['piece_mvt'], r['opus'], r['instrument'], r['style'], r['date'],
            r['measures'], r['m_method'], r['pages'], r['p_method'],
            r['size_kb'], r['folder_size_kb'], r['ily_count'], r['rel_path'],
            r['filename'], r['version'], r['license'], r['maintainer']
        ]
        ws1.append(row_values)
        ws1.row_dimensions[r_idx].height = 20

        is_even = (r_idx % 2 == 0)
        for col_idx in range(1, len(row_values) + 1):
            cell = ws1.cell(row=r_idx, column=col_idx)
            cell.font = font_data
            cell.border = border_thin
            if is_even:
                cell.fill = fill_zebra

            # Alignments
            if col_idx in [1, 2, 4, 10, 11, 13, 17]:
                cell.alignment = align_center
            elif col_idx in [15, 16]:
                cell.alignment = align_right
                cell.number_format = '#,##0.0'
            else:
                cell.alignment = align_left

            if col_idx in [1, 11, 13, 17]:
                cell.number_format = '#,##0'

    # Sheet 2: Composer Summary
    ws2 = wb.create_sheet(title="작곡가별 집계")
    ws2.views.sheetView[0].showGridLines = True

    comp_headers = [
        "순위", "작곡가", "작곡가 코드", "보유 악보 수 (곡/악장)", "총 마디 수 합계",
        "총 예상 쪽수 합계", "평균 곡당 마디", "평균 곡당 쪽수", "총 파일 용량 (MB)",
        "주요 음악 양식", "주요 악기 / 편성"
    ]
    ws2.append(comp_headers)
    ws2.row_dimensions[1].height = 28

    fill_comp_header = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
    for col_num in range(1, len(comp_headers) + 1):
        cell = ws2.cell(row=1, column=col_num)
        cell.font = font_header
        cell.fill = fill_comp_header
        cell.alignment = align_center

    sorted_composers = sorted(composer_stats.items(), key=lambda x: x[1]['count'], reverse=True)

    for rank, (comp_code, cs) in enumerate(sorted_composers, start=1):
        r_idx = rank + 1
        display_name = comp_code
        if cs['names']:
            display_name = max(cs['names'].keys(), key=lambda k: (cs['names'][k], len(k)))

        avg_m = round(cs['total_measures'] / cs['count'], 1) if cs['count'] else 0
        avg_p = round(cs['total_pages'] / cs['count'], 1) if cs['count'] else 0
        mb_size = round(cs['total_size_kb'] / 1024, 2)
        inst_str = ', '.join(sorted(list(cs['instruments']))[:4])
        style_str = ', '.join(sorted(list(cs['styles'])))

        row_vals = [
            rank, display_name, cs['code'], cs['count'], cs['total_measures'],
            cs['total_pages'], avg_m, avg_p, mb_size, style_str, inst_str
        ]
        ws2.append(row_vals)
        ws2.row_dimensions[r_idx].height = 20
        is_even = (r_idx % 2 == 0)

        for col_idx in range(1, len(row_vals) + 1):
            cell = ws2.cell(row=r_idx, column=col_idx)
            cell.font = font_data
            cell.border = border_thin
            if is_even:
                cell.fill = fill_zebra

            if col_idx in [1, 3]:
                cell.alignment = align_center
                cell.number_format = '#,##0'
            elif col_idx in [4, 5, 6]:
                cell.alignment = align_right
                cell.number_format = '#,##0'
            elif col_idx in [7, 8, 9]:
                cell.alignment = align_right
                cell.number_format = '#,##0.0'
            else:
                cell.alignment = align_left

    # Auto-adjust column widths for both sheets
    for ws in [ws1, ws2]:
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col[:100]:  # sample top 100 rows for speed
                val = str(cell.value or '')
                # Approximate East Asian width
                length = sum(2 if ord(c) > 127 else 1 for c in val)
                if length > max_len:
                    max_len = length
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # Save to both paths
    print(f"Saving workbook to {OUTPUT_XLSX} and {OUTPUT_XLSX_ALT}...")
    wb.save(OUTPUT_XLSX)
    wb.save(OUTPUT_XLSX_ALT)

    print("\n==========================================")
    print("Mutopia Excel catalog generation complete!")
    print(f"Total scores cataloged: {len(records)}")
    print(f"Total composers summarized: {len(sorted_composers)}")
    print(f"Saved to: {OUTPUT_XLSX}")
    print(f"File size: {round(os.path.getsize(OUTPUT_XLSX) / 1024, 1)} KB")

if __name__ == '__main__':
    main()
