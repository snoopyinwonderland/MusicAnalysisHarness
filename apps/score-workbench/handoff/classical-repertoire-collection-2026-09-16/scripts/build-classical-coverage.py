import json
import re
import sqlite3
import sys
import unicodedata
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

workbook_path = Path(sys.argv[1])
db_path = Path(sys.argv[2])
output_path = Path(sys.argv[3])

def norm(value):
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(ch for ch in text if not unicodedata.combining(ch)).lower()
    text = text.replace("ß", "ss").replace("œ", "oe").replace("æ", "ae")
    text = re.sub(r"\b(?:composed|composer|music|lyrics|arranged|arranger|by)\b", " ", text)
    return " ".join(re.findall(r"[a-z0-9]+", text))

def tokens(value):
    stop = {"the","a","an","in","on","of","for","and","no","number","major","minor","op","opus","movement","mov","part","book"}
    return {t for t in norm(value).split() if t not in stop and len(t) > 1}

def catalogue(value):
    patterns = [r"\bBWV\s*Anh\.?\s*\d+[a-z]?", r"\bBWV\s*\d+[a-z]?", r"\bK(?:V)?\.?\s*\d+[a-z]?", r"\bHob\.?\s*[IVX]+[: .-]*\d+", r"\bRV\s*\d+", r"\bHWV\s*\d+", r"\bD\.?\s*\d+[a-z]?", r"\bS\.?\s*\d+[a-z]?", r"\bWWV\s*\d+", r"\bWAB\s*\d+", r"\bB\.?\s*\d+", r"\bL\.?\s*\d+", r"\bM\.?\s*\d+", r"\bOp\.?\s*\d+(?:\s*No\.?\s*\d+)?"]
    for pattern in patterns:
        m = re.search(pattern, str(value or ""), re.I)
        if m:
            return re.sub(r"[^A-Z0-9]", "", m.group(0).upper())
    return ""

def composer_key(value):
    parts = norm(value).split()
    return parts[-1] if parts else ""

wb = load_workbook(workbook_path, read_only=True, data_only=True)
sheet = wb["작품 목록"]
targets = []
for row in sheet.iter_rows(min_row=3, values_only=True):
    if not row[6] or not row[7]:
        continue
    targets.append({
        "row": len(targets) + 3,
        "tier": row[1], "score": row[2], "era": row[3], "composer_grade": row[5],
        "composer": row[6], "title": row[7], "catalogue": row[12] or catalogue(row[7]),
    })
wb.close()

db = sqlite3.connect(f"file:{db_path.as_posix()}?mode=ro", uri=True)
db.row_factory = sqlite3.Row
rows = db.execute("select id,title,normalized_title,composer,source,license,access_policy,pdmx_metadata,source_metadata from works").fetchall()
stream_count = len(rows)
grouped = {}
for row in rows:
    key = (row["source"] or row["id"], row["title"] or "", row["composer"] or "")
    if key in grouped:
        grouped[key]["streams"] += 1
        continue
    metadata = " ".join([row["pdmx_metadata"] or "", row["source_metadata"] or ""])
    grouped[key] = {
        "id": row["id"], "title": row["title"] or "", "composer": row["composer"] or "",
        "source": row["source"] or "", "license": row["license"] or "",
        "access_policy": row["access_policy"] or "", "metadata": metadata, "streams": 1,
    }
db.close()

inventory = list(grouped.values())
for item in inventory:
    item["title_norm"] = norm(item["title"])
    item["title_tokens"] = tokens(item["title"])
    item["catalogue"] = catalogue(" ".join([item["title"], item["source"], item["metadata"]]))
    item["composer_norm"] = norm(" ".join([item["composer"], item["metadata"]]))
    source_lower = item["source"].lower().replace("\\", "/")
    item["dataset"] = "PDMX" if source_lower.startswith("mxl/") or '"metadata_path"' in item["metadata"] else "KYSing" if source_lower.startswith("kysing/") or '"catalog":"kumyoung"' in item["metadata"].replace(" ", "") else "Music Analysis"

catalog_index = defaultdict(list)
title_index = defaultdict(list)
for item in inventory:
    if item["catalogue"]:
        catalog_index[item["catalogue"]].append(item)
    for tok in item["title_tokens"]:
        title_index[tok].append(item)

matches = []
for target in targets:
    tcat = re.sub(r"[^A-Z0-9]", "", str(target["catalogue"]).upper()) if target["catalogue"] else ""
    ttok = tokens(target["title"])
    tnorm = norm(target["title"])
    surname = composer_key(target["composer"])
    candidates = list(catalog_index.get(tcat, [])) if tcat else []
    if not candidates and ttok:
        rare = sorted(ttok, key=lambda t: len(title_index.get(t, [])))[:3]
        seen = set()
        for tok in rare:
            for item in title_index.get(tok, []):
                key = item["source"] or item["id"]
                if key not in seen:
                    candidates.append(item); seen.add(key)
    best = None
    for item in candidates:
        overlap = len(ttok & item["title_tokens"]) / max(1, len(ttok | item["title_tokens"]))
        containment = len(ttok & item["title_tokens"]) / max(1, min(len(ttok), len(item["title_tokens"])))
        exact_title = bool(tnorm and (tnorm == item["title_norm"] or tnorm in item["title_norm"] or item["title_norm"] in tnorm))
        catalog_match = bool(tcat and tcat == item["catalogue"])
        composer_match = bool(surname and surname in item["composer_norm"])
        score = (0.62 if catalog_match else 0) + (0.25 * overlap) + (0.18 * containment) + (0.18 if exact_title else 0) + (0.15 if composer_match else 0)
        if catalog_match and not composer_match:
            score -= 0.12
        if best is None or score > best[0]:
            best = (score, item, catalog_match, exact_title, composer_match, overlap)
    status = "미보유"
    confidence = 0.0
    method = "일치 없음"
    matched = None
    if best:
        confidence, matched, cat_ok, title_ok, comp_ok, overlap = best
        if cat_ok and (comp_ok or title_ok): status, method = "보유", "카탈로그 번호+작곡가/제목"
        elif title_ok and comp_ok: status, method = "보유", "제목+작곡가"
        elif comp_ok and overlap >= 0.62: status, method = "부분 보유", "작곡가+제목 토큰"
        elif cat_ok or (title_ok and overlap >= 0.5): status, method = "검토 필요", "단일 강한 단서"
        else: matched = None
    matches.append({
        **target, "status": status, "confidence": round(min(1, max(0, confidence)), 4), "method": method,
        "matched_title": matched["title"] if matched else "", "matched_composer": matched["composer"] if matched else "",
        "source": matched["source"] if matched else "", "work_id": matched["id"] if matched else "",
        "dataset": matched["dataset"] if matched else "", "streams": matched["streams"] if matched else 0,
        "license": matched["license"] if matched else "", "access_policy": matched["access_policy"] if matched else "",
    })

summary = defaultdict(int)
for match in matches:
    summary[match["status"]] += 1
dataset_summary = defaultdict(int)
for item in inventory:
    dataset_summary[item["dataset"]] += 1

payload = {
    "schema": "musicanote-classical-coverage-v1", "workbook": str(workbook_path), "database": str(db_path),
    "stream_count": stream_count, "inventory_works": len(inventory), "dataset_summary": dict(dataset_summary),
    "target_works": len(targets), "summary": dict(summary), "matches": matches,
    "methodology": {
        "automatic_owned": "카탈로그 번호+작곡가/제목 또는 정규화 제목+작곡가가 함께 일치",
        "partial": "작곡가와 제목 핵심 토큰 62% 이상",
        "review": "카탈로그 또는 제목 단일 단서만 강한 경우",
        "unowned": "위 기준 미충족",
    },
}
output_path.parent.mkdir(parents=True, exist_ok=True)
output_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps({k: payload[k] for k in ["stream_count","inventory_works","dataset_summary","target_works","summary"]}, ensure_ascii=False, indent=2))
