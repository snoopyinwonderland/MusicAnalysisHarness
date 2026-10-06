"""Create a non-destructive authority-list collection workspace from v4.

This script never writes to the source workbook and refuses to replace an
existing generated artifact.  It turns the handoff workbook's pending list
collection work into versioned JSONL queues.
"""

from __future__ import annotations

import argparse
import hashlib
import json
from collections import Counter
from datetime import date
from pathlib import Path

import openpyxl


RUN_DATE = date(2026, 9, 22).isoformat()
SCHEMA = "musicanote-classical-authority-backlog-v1"


def stable_id(prefix: str, value: str) -> str:
    digest = hashlib.sha256(value.strip().casefold().encode("utf-8")).hexdigest()
    return f"{prefix}_{digest[:16]}"


def row_records(sheet):
    headers = [sheet.cell(2, column).value for column in range(1, sheet.max_column + 1)]
    for row_number in range(3, sheet.max_row + 1):
        values = [sheet.cell(row_number, column).value for column in range(1, sheet.max_column + 1)]
        if not any(value is not None for value in values):
            continue
        yield row_number, dict(zip(headers, values))


def write_jsonl(path: Path, records: list[dict]) -> None:
    if path.exists():
        raise FileExistsError(f"Refusing to overwrite existing artifact: {path}")
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("x", encoding="utf-8", newline="\n") as handle:
        for record in records:
            handle.write(json.dumps(record, ensure_ascii=False, sort_keys=True))
            handle.write("\n")


def write_text(path: Path, contents: str) -> None:
    if path.exists():
        raise FileExistsError(f"Refusing to overwrite existing artifact: {path}")
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("x", encoding="utf-8", newline="\n") as handle:
        handle.write(contents)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source_workbook", type=Path)
    parser.add_argument("output_root", type=Path)
    args = parser.parse_args()

    source = args.source_workbook.resolve()
    output_root = args.output_root.resolve()
    if not source.is_file():
        raise FileNotFoundError(source)
    if output_root.exists():
        raise FileExistsError(f"Refusing to reuse existing workspace: {output_root}")

    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    composer_rows = list(row_records(workbook["작곡가 목록"]))
    catalogue_rows = list(row_records(workbook["전작 카탈로그"]))
    supplementary_rows = list(row_records(workbook["보완 작곡가"]))
    works_rows = list(row_records(workbook["작품 목록"]))

    composer_records = []
    composer_ids = {}
    policy_counts = Counter()
    complete_status_counts = Counter()
    year_status_counts = Counter()
    for source_row, item in composer_rows:
        canonical = item["작곡가"]
        composer_id = stable_id("cmp", canonical)
        composer_ids[canonical] = composer_id
        policy = item["수집 정책"]
        complete_status = item["Complete Catalogue 상태"]
        year_status = item["작품 연도 보완"]
        policy_counts[policy] += 1
        complete_status_counts[complete_status] += 1
        year_status_counts[year_status] += 1
        composer_records.append(
            {
                "schemaVersion": SCHEMA,
                "composerId": composer_id,
                "composerCanonical": canonical,
                "era": item["시대"],
                "birthYear": item["출생 연도"],
                "deathYear": item["사망 연도"],
                "priorityGrade": item["작곡가 등급"],
                "collectionPolicy": policy,
                "completeCatalogueStatus": complete_status,
                "workDateEnrichmentStatus": year_status,
                "openOpusId": item["Open Opus ID"],
                "legacySourceUrl": item["출처 URL"],
                "authorityCollectionStatus": "waiting",
                "legacyWorkbookRow": source_row,
            }
        )

    catalogue_records = []
    for source_row, item in catalogue_rows:
        composer = item["작곡가"]
        catalogue_records.append(
            {
                "schemaVersion": SCHEMA,
                "composerId": composer_ids.get(composer, stable_id("cmp", composer)),
                "composerCanonical": composer,
                "priorityGrade": item["등급"],
                "catalogueAbbreviation": item["약호"],
                "catalogueCanonicalName": item["카탈로그"],
                "authorityProvider": item["우선 출처"],
                "authorityUrl": item["URL"],
                "collectionNotes": item["사용 메모"],
                "handoffVerificationStatus": item["확인 상태"],
                "internalMappingStatus": item["내부 매핑 상태"],
                "authorityCollectionStatus": "waiting",
                "legacyWorkbookRow": source_row,
            }
        )

    supplementary_records = []
    for source_row, item in supplementary_rows:
        canonical = item["작곡가"]
        supplementary_records.append(
            {
                "schemaVersion": SCHEMA,
                "composerId": stable_id("cmp", canonical),
                "composerCanonical": canonical,
                "era": item["시대"],
                "birthYear": item["출생"],
                "deathYear": item["사망"],
                "priorityGrade": item["제안 등급"],
                "region": item["지역"],
                "inclusionRationale": item["보완 목적"],
                "representativeWorks": item["대표 작품"],
                "initialRightsClassification": item["권리 1차 분류"],
                "verificationNotes": item["검증 메모"],
                "authorityCollectionStatus": "waiting",
                "legacyWorkbookRow": source_row,
            }
        )

    work_year_missing = sum(1 for _, item in works_rows if item["작곡 연도"] is None)
    work_catalogue_missing = sum(1 for _, item in works_rows if item["카탈로그 번호"] is None)
    expected_paths = [
        output_root / "README.md",
        output_root / "catalog" / "authority" / "composer-backlog.jsonl",
        output_root / "catalog" / "authority" / "s-composer-catalogue-registry.jsonl",
        output_root / "catalog" / "authority" / "supplementary-composer-backlog.jsonl",
        output_root / "reports" / f"list-collection-audit-{RUN_DATE}.md",
        output_root / "changes" / "CHANGELOG.md",
        output_root / "changes" / "events.jsonl",
    ]
    if any(path.exists() for path in expected_paths):
        existing = next(path for path in expected_paths if path.exists())
        raise FileExistsError(f"Refusing to overwrite existing artifact: {existing}")

    write_jsonl(output_root / "catalog" / "authority" / "composer-backlog.jsonl", composer_records)
    write_jsonl(
        output_root / "catalog" / "authority" / "s-composer-catalogue-registry.jsonl",
        catalogue_records,
    )
    write_jsonl(
        output_root / "catalog" / "authority" / "supplementary-composer-backlog.jsonl",
        supplementary_records,
    )
    write_jsonl(
        output_root / "changes" / "events.jsonl",
        [
            {
                "eventId": "evt_20260922_bootstrap_authority_backlog",
                "eventDate": RUN_DATE,
                "eventType": "workspace-bootstrap",
                "sourceWorkbook": str(source),
                "sourceWorkbookSha256": hashlib.sha256(source.read_bytes()).hexdigest(),
                "outputs": [str(path.relative_to(output_root)) for path in expected_paths[1:]],
                "notes": "Generated pending authority-list queues without modifying the source workbook.",
            }
        ],
    )
    write_text(
        output_root / "README.md",
        "# Classical repertoire collection\n\n"
        "This workspace is a non-destructive continuation of the 2026-09-16 handoff. "
        "The v4 workbook remains the immutable baseline. JSONL files under `catalog/authority/` "
        "are pending queues, not completed authority records.\n",
    )
    write_text(
        output_root / "changes" / "CHANGELOG.md",
        "# Change log\n\n"
        f"## {RUN_DATE}\n\n"
        "- Created authority-list backlog files from the immutable v4 workbook.\n"
        "- Did not modify the source workbook or mark any pending item as verified.\n",
    )
    audit = f"""# List collection audit — {RUN_DATE}

## Baseline

- Source workbook: `{source}`
- Composers in baseline: {len(composer_records)}
- Works in baseline: {len(works_rows)}
- S-grade catalogue registry rows: {len(catalogue_records)}
- Supplementary-composer candidates: {len(supplementary_records)}

## Pending authority work

- Complete Catalogue statuses: `{dict(sorted(complete_status_counts.items()))}`
- Work-year enrichment statuses: `{dict(sorted(year_status_counts.items()))}`
- Collection policies: `{dict(sorted(policy_counts.items()))}`
- Works without composition year: {work_year_missing}
- Works without catalogue number: {work_catalogue_missing}

## Interpretation

The baseline is a broad discovery list. It is not yet a completed authority catalogue:
the authority queues above must be resolved from composer-specific catalogues or other
documented sources. A pending queue must never be treated as a verified work list.
"""
    write_text(output_root / "reports" / f"list-collection-audit-{RUN_DATE}.md", audit)
    print(f"Created {output_root}")
    print(f"composer backlog: {len(composer_records)}")
    print(f"S-grade catalogue registry: {len(catalogue_records)}")
    print(f"supplementary composer backlog: {len(supplementary_records)}")


if __name__ == "__main__":
    main()
