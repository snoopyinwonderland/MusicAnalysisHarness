"""Run the backlog bootstrap with reliable worksheet dimensions."""

from __future__ import annotations

import sys

import openpyxl

from bootstrap_authority_backlog import main


def load_workbook_with_dimensions(*args, **kwargs):
    kwargs["read_only"] = False
    kwargs["data_only"] = True
    return original_load_workbook(*args, **kwargs)


original_load_workbook = openpyxl.load_workbook
openpyxl.load_workbook = load_workbook_with_dimensions

sys.argv = [
    "bootstrap_authority_backlog.py",
    r"K:\MusicSearch\outputs\classical-repertoire-priority-2026-09-08\classical-repertoire-priority-v4.xlsx",
    r"K:\MusicSearch\classical-repertoire-collection\workspace",
]
main()
